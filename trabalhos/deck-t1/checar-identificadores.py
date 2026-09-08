# -*- coding: utf-8 -*-
"""Varre o .pptx atras de identificador interno que a plateia nao consegue resolver.
Uso: python checar-identificadores.py <arquivo.pptx>
Sai com codigo 1 se achar alguma coisa."""
import re, sys
from pptx import Presentation

# o que nao pode aparecer: identificador que so existe no board ou no plano de ensino do grupo
PROIBIDO = [
    ('caso de uso',        r'\bUC\s?\d+\b'),
    ('work item',          r'#\d+'),
    ('trabalho da cadeira', r'\bT[1-4]\b'),
    ('aula',               r'\bAula\s?0?\d\b'),
    # nome de artefato interno, que a plateia nao tem como resolver
    ('artefato do grupo',  r'\bwiki\b|Guia da Arquitetura|Team Charter|\bTo Do\b|\bDoing\b|\bbacklog\b|\bMarina\b|\bRafael\b|\bJorge\b'),
    ('sprint numerada',    r'\bSprint\s?\d'),
    ('jargao do metodo',   r'Quest[ãa]o Norteadora|\bADR\b'),
]
# onde a mencao e legitima: a capa identifica o trabalho, a pagina de referencias cita a bibliografia.
# A capa e sempre o slide 1. A pagina de referencias e reconhecida pelo titulo, para nao depender do numero.
ISENTO_POR_SLIDE = {'trabalho da cadeira': {1}}
ISENTO_POR_TITULO = {'aula': 'referências'}

def main(caminho):
    prs = Presentation(caminho)
    achados = []
    for i, sl in enumerate(prs.slides, 1):
        txt = ' '.join(' '.join(sh.text_frame.text.split())
                       for sh in sl.shapes if sh.has_text_frame)
        for nome, pat in PROIBIDO:
            if i in ISENTO_POR_SLIDE.get(nome, set()):
                continue
            if nome in ISENTO_POR_TITULO and ISENTO_POR_TITULO[nome] in txt.lower()[:60]:
                continue
            for m in re.finditer(pat, txt):
                ini, fim = max(0, m.start() - 50), min(len(txt), m.end() + 30)
                achados.append((i, nome, m.group(0), txt[ini:fim]))
    for i, nome, tok, ctx in achados:
        print(f'slide {i:02d}  {nome}  [{tok}]  ...{ctx}...')
    print(f'{len(achados)} identificador(es) sem contexto' if achados else 'nenhum identificador solto')
    return 1 if achados else 0

if __name__ == '__main__':
    sys.exit(main(sys.argv[1]))
