# -*- coding: utf-8 -*-
"""Varre o .pptx atras de identificador interno que a plateia nao consegue resolver.
Uso: python checar-identificadores.py <arquivo.pptx>
Sai com codigo 1 se achar alguma coisa."""
import re, sys
from pptx import Presentation


PROIBIDO = [
    ('caso de uso',        r'\bUC\s?\d+\b'),
    ('work item',          r'#\d+'),
    ('trabalho da cadeira', r'\bT[1-4]\b'),
    ('jargao do metodo',   r'Quest[ãa]o Norteadora|\bADR\b'),
    ('linha do guia',      r'\blinha\s\d+'),
    ('nome de arquivo',    r'\b[\w-]+\.(ts|tsx|mjs|cjs|json|yml|yaml|sql|md|py)\b'),
    ('estado do quadro',   r'\bTo Do\b|\bDoing\b'),
    ('nome de persona',    r'\bMarina\b|\bRafael\b|\bJorge\b'),
    ('codigo de historia', r'\bH\d{2}\b'),
    ('nome provisorio',    r'Invite[ -]?People|invite-app'),
    ('ponto medio',        r'·'),
]
# As sprints ficam no slide, porque sao as do cronograma da disciplina. Os codigos das historias e o
# nome provisorio do projeto ficam so na fala. A capa identifica o trabalho.
ISENTO_POR_SLIDE = {'trabalho da cadeira': {1}}
ISENTO_POR_TITULO = {}

def main(caminho):
    prs = Presentation(caminho)
    achados = []
    for i, sl in enumerate(prs.slides, 1):
        txt = ' '.join(' '.join(sh.text_frame.text.split())
                       for sh in sl.shapes if sh.has_text_frame)
        for nome, pat in PROIBIDO:
            if i in ISENTO_POR_SLIDE.get(nome, set()):
                continue
            if nome in ISENTO_POR_TITULO and ISENTO_POR_TITULO[nome] in txt.lower()[:60]:
                continue
            for m in re.finditer(pat, txt):
                ini, fim = max(0, m.start() - 50), min(len(txt), m.end() + 30)
                achados.append((i, nome, m.group(0), txt[ini:fim]))
    for i, nome, tok, ctx in achados:
        print(f'slide {i:02d}  {nome}  [{tok}]  ...{ctx}...')
    print(f'{len(achados)} identificador(es) sem contexto' if achados else 'nenhum identificador solto')
    return 1 if achados else 0

if __name__ == '__main__':
    sys.exit(main(sys.argv[1]))
